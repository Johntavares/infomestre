from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
import sys
import os

def fix_presentation(file_path, out_path):
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text_frame = shape.text_frame
            # Enable auto fit text to shape and word wrap
            text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            text_frame.word_wrap = True
            
            # For slide 8 specifically, we might also want to increase shape height just in case
            if i + 1 == 8:
                try:
                    # increase height by 20%
                    shape.height = int(shape.height * 1.2)
                except:
                    pass

    prs.save(out_path)
    print(f"Fixed presentation saved to {out_path}")

if __name__ == '__main__':
    in_file = sys.argv[1]
    out_file = sys.argv[2]
    fix_presentation(in_file, out_file)
